# -*- coding: utf-8 -*-
"""
Generates the academic project report (ספר הפרויקט) for TrendAnalyzer Pro as a
Word document, following the Tel-Hai Academic College formatting conventions:
right-to-left Hebrew body in David, 1.5 line spacing, numbered chapters, a table of
contents and lists of figures and tables built from real Word fields.

Run from the project root:

    venv/bin/python docs/generate_report.py

Output: TrendAnalyzer_Pro_Project_Report.docx in the project root. Word fills in
page numbers for the table of contents and the lists of figures/tables when the
document is opened (or with F9 / "Update Field").
"""

import os
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------------------------------------------------------------------------
# Document settings
# ---------------------------------------------------------------------------
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUTPUT_PATH = os.path.join(PROJECT_ROOT, "TrendAnalyzer_Pro_Project_Report.docx")

HEB_FONT = "David"
LATIN_FONT = "David"
CODE_FONT = "Consolas"
EN_FONT = "Times New Roman"
BODY_PT = 12
LINE_SPACING = 1.5

NAVY = RGBColor(0x1F, 0x2A, 0x44)
ACCENT = RGBColor(0x2A, 0x78, 0xD6)
HEADER_FILL = "1F2A44"
BAND_FILLS = ("DCE8F7", "E3F4EC", "F6E6DC", "ECE8F6", "F3F4F6")

FIG_LABEL = "איור"
TAB_LABEL = "טבלה"

# Latin segments inside Hebrew text are written as separate left-to-right runs so
# that Word's bidi algorithm never reorders English terms, numbers or formulas.
_LATIN_RE = re.compile(r"[A-Za-z][A-Za-z0-9 .,_\-+/:#=<>≥≤×·'\"%&*\[\]]*[A-Za-z0-9%\]]|[A-Za-z]")


# ---------------------------------------------------------------------------
# Low-level OOXML helpers
# ---------------------------------------------------------------------------
def _insert_before_any(parent, child, local_names):
    """Inserts `child` before the first existing child whose local name is listed,
    preserving the schema order Word validates; appends when none is present."""
    for existing in parent:
        if existing.tag.split('}')[-1] in local_names:
            existing.addprevious(child)
            return child
    parent.append(child)
    return child


_PPR_AFTER_BIDI = ("adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
                   "mirrorIndents", "suppressOverlap", "jc", "textDirection",
                   "textAlignment", "textboxTightWrap", "outlineLvl", "divId",
                   "cnfStyle", "rPr", "sectPr", "pPrChange")
_RPR_AFTER_RFONTS = ("b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike",
                     "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid",
                     "vanish", "webHidden", "color", "spacing", "w", "kern", "position",
                     "sz", "szCs", "highlight", "u", "effect", "bdr", "shd", "fitText",
                     "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout",
                     "specVanish", "oMath")


def _set_bidi(paragraph_or_style_ppr):
    ppr = paragraph_or_style_ppr
    if ppr.find(qn('w:bidi')) is None:
        _insert_before_any(ppr, OxmlElement('w:bidi'), _PPR_AFTER_BIDI)


def _rpr_child(rpr, tag, after):
    el = rpr.find(qn('w:' + tag))
    if el is None:
        el = _insert_before_any(rpr, OxmlElement('w:' + tag), after)
    return el


def _set_fonts(rpr, latin, cs):
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts) if rpr.find(qn('w:rStyle')) is None else \
            rpr.find(qn('w:rStyle')).addnext(fonts)
    for attr in list(fonts.attrib):
        if attr.endswith('Theme'):
            del fonts.attrib[attr]
    fonts.set(qn('w:ascii'), latin)
    fonts.set(qn('w:hAnsi'), latin)
    fonts.set(qn('w:eastAsia'), latin)
    fonts.set(qn('w:cs'), cs)


def _set_cs_size(rpr, pt):
    idx = _RPR_AFTER_RFONTS.index("szCs")
    sz = _rpr_child(rpr, "sz", _RPR_AFTER_RFONTS[_RPR_AFTER_RFONTS.index("sz") + 1:])
    sz.set(qn('w:val'), str(int(pt * 2)))
    szcs = _rpr_child(rpr, "szCs", _RPR_AFTER_RFONTS[idx + 1:])
    szcs.set(qn('w:val'), str(int(pt * 2)))


def _set_flag(rpr, tag, on=True):
    idx = _RPR_AFTER_RFONTS.index(tag)
    el = _rpr_child(rpr, tag, _RPR_AFTER_RFONTS[idx + 1:])
    if not on:
        el.set(qn('w:val'), '0')


def _set_lang(rpr):
    lang = _rpr_child(rpr, "lang", _RPR_AFTER_RFONTS[_RPR_AFTER_RFONTS.index("lang") + 1:])
    lang.set(qn('w:val'), 'en-US')
    lang.set(qn('w:bidi'), 'he-IL')


def _shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    _insert_before_any(tcpr, shd, ("noWrap", "tcMar", "textDirection", "tcFitText",
                                   "vAlign", "hideMark", "headers", "cellIns",
                                   "cellDel", "cellMerge", "tcPrChange"))


def _cell_borders(cell, style="single", size=8, color="1F2A44"):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), style)
        el.set(qn('w:sz'), str(size))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    _insert_before_any(tcpr, borders, ("shd", "noWrap", "tcMar", "textDirection",
                                       "tcFitText", "vAlign", "hideMark"))


def _table_bidi(table):
    tblpr = table._tbl.tblPr
    _insert_before_any(tblpr, OxmlElement('w:bidiVisual'),
                       ("tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd",
                        "tblLayout", "tblCellMar", "tblLook", "tblCaption",
                        "tblDescription"))


def _repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement('w:tblHeader'))


def _keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


# ---------------------------------------------------------------------------
# Field helpers (TOC, SEQ, PAGE)
# ---------------------------------------------------------------------------
def _fld_char(run, kind, dirty=False):
    fc = OxmlElement('w:fldChar')
    fc.set(qn('w:fldCharType'), kind)
    if dirty:
        fc.set(qn('w:dirty'), 'true')
    run._r.append(fc)


def _instr(run, text):
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = f' {text} '
    run._r.append(it)


def add_simple_field(paragraph, instruction, result_text, rtl=True):
    """A complete field inside one paragraph, with a pre-computed result."""
    r = paragraph.add_run()
    _fld_char(r, 'begin')
    r = paragraph.add_run()
    _instr(r, instruction)
    r = paragraph.add_run()
    _fld_char(r, 'separate')
    r = paragraph.add_run(result_text)
    _style_run(r, rtl=rtl)
    r = paragraph.add_run()
    _fld_char(r, 'end')


# ---------------------------------------------------------------------------
# Document builder
# ---------------------------------------------------------------------------
def _style_run(run, rtl=True, bold=None, italic=None, size=None, font=None, color=None):
    rpr = run._r.get_or_add_rPr()
    if font:
        _set_fonts(rpr, font, font)
    if bold:
        _set_flag(rpr, "b")
        _set_flag(rpr, "bCs")
    if italic:
        _set_flag(rpr, "i")
        _set_flag(rpr, "iCs")
    if size:
        _set_cs_size(rpr, size)
    if color is not None:
        run.font.color.rgb = color
    if rtl:
        _set_flag(rpr, "rtl")


class Report:
    """Thin wrapper around python-docx that knows about RTL Hebrew and captions."""

    def __init__(self):
        self.doc = Document()
        self.fig_no = 0
        self.tab_no = 0
        self.headings = []      # (level, text) for the pre-rendered TOC
        self.figures = []       # caption texts
        self.tables = []
        self._setup_page()
        self._setup_styles()

    # -- setup -------------------------------------------------------------
    def _setup_page(self):
        section = self.doc.sections[0]
        section.page_height, section.page_width = Cm(29.7), Cm(21.0)
        section.orientation = WD_ORIENT.PORTRAIT
        section.top_margin = section.bottom_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.different_first_page_header_footer = True
        # Right-to-left section (page and column flow).
        sectpr = section._sectPr
        _insert_before_any(sectpr, OxmlElement('w:bidi'),
                           ("rtlGutter", "docGrid", "printerSettings", "sectPrChange"))
        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_bidi(footer_p._p.get_or_add_pPr())
        add_simple_field(footer_p, 'PAGE', '1', rtl=False)

        header_p = section.header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_bidi(header_p._p.get_or_add_pPr())
        r = header_p.add_run("TrendAnalyzer Pro — ספר פרויקט")
        _style_run(r, size=9, color=RGBColor(0x6B, 0x72, 0x80))

        # Ask Word to refresh the TOC / lists on open.
        settings = self.doc.settings.element
        upd = OxmlElement('w:updateFields')
        upd.set(qn('w:val'), 'true')
        _insert_before_any(settings, upd, (
            "hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids",
            "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping",
            "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
            "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
            "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator"))

    def _setup_styles(self):
        styles = self.doc.styles
        normal = styles['Normal']
        rpr = normal.element.get_or_add_rPr()
        _set_fonts(rpr, LATIN_FONT, HEB_FONT)
        _set_cs_size(rpr, BODY_PT)
        _set_lang(rpr)
        pf = normal.paragraph_format
        pf.line_spacing = LINE_SPACING
        pf.space_after = Pt(6)
        _set_bidi(normal.element.get_or_add_pPr())

        for level, size in ((1, 18), (2, 15), (3, 13)):
            st = styles[f'Heading {level}']
            rpr = st.element.get_or_add_rPr()
            _set_fonts(rpr, HEB_FONT, HEB_FONT)
            _set_cs_size(rpr, size)
            _set_flag(rpr, "b")
            _set_flag(rpr, "bCs")
            st.font.color.rgb = NAVY if level == 1 else ACCENT if level == 2 else NAVY
            st.paragraph_format.space_before = Pt(18 if level == 1 else 12)
            st.paragraph_format.space_after = Pt(8)
            st.paragraph_format.line_spacing = 1.15
            st.paragraph_format.keep_with_next = True
            _set_bidi(st.element.get_or_add_pPr())

        try:
            cap = styles['Caption']
        except KeyError:
            from docx.enum.style import WD_STYLE_TYPE
            cap = styles.add_style('Caption', WD_STYLE_TYPE.PARAGRAPH)
            cap.base_style = normal
        rpr = cap.element.get_or_add_rPr()
        _set_fonts(rpr, HEB_FONT, HEB_FONT)
        _set_cs_size(rpr, 10)
        _set_flag(rpr, "b")
        _set_flag(rpr, "bCs")
        cap.font.italic = False
        cap.font.color.rgb = NAVY
        cap.paragraph_format.space_after = Pt(12)
        cap.paragraph_format.line_spacing = 1.0
        _set_bidi(cap.element.get_or_add_pPr())

        for name in ('TOC 1', 'TOC 2', 'TOC 3', 'Table of Figures'):
            try:
                st = styles[name]
            except KeyError:
                continue
            _set_bidi(st.element.get_or_add_pPr())

    # -- text --------------------------------------------------------------
    def _add_mixed_text(self, p, text, bold=False, italic=False, size=None, color=None):
        """Adds Hebrew text with embedded Latin segments as correctly-directed runs.
        `**bold**` spans inside `text` are honoured."""
        for i, chunk in enumerate(re.split(r'\*\*', text)):
            chunk_bold = bold or (i % 2 == 1)
            pos = 0
            for m in _LATIN_RE.finditer(chunk):
                if m.start() > pos:
                    r = p.add_run(chunk[pos:m.start()])
                    _style_run(r, rtl=True, bold=chunk_bold, italic=italic, size=size,
                               color=color)
                r = p.add_run(m.group(0))
                _style_run(r, rtl=False, bold=chunk_bold, italic=italic, size=size,
                           color=color)
                pos = m.end()
            if pos < len(chunk):
                r = p.add_run(chunk[pos:])
                _style_run(r, rtl=True, bold=chunk_bold, italic=italic, size=size,
                           color=color)

    def para(self, text, align='justify', bold=False, size=None, color=None,
             space_after=None, style=None):
        p = self.doc.add_paragraph(style=style)
        _set_bidi(p._p.get_or_add_pPr())
        if align == 'justify':
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif align == 'center':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        self._add_mixed_text(p, text, bold=bold, size=size, color=color)
        return p

    def bullets(self, items, numbered=False):
        for idx, item in enumerate(items, 1):
            p = self.doc.add_paragraph()
            _set_bidi(p._p.get_or_add_pPr())
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.left_indent = Cm(0.9)  # start side in an RTL paragraph
            pf.first_line_indent = Cm(-0.6)
            pf.space_after = Pt(3)
            marker = f"{idx}. " if numbered else "• "
            r = p.add_run(marker)
            _style_run(r, rtl=True, bold=numbered)
            self._add_mixed_text(p, item)

    def heading(self, text, level=1, page_break=False):
        if page_break:
            self.page_break()
        h = self.doc.add_heading(level=level)
        _set_bidi(h._p.get_or_add_pPr())
        self._add_mixed_text(h, text)
        self.headings.append((level, text))
        return h

    def page_break(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    def code(self, text):
        """Left-to-right monospace block (code or formula), on a light background."""
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        _shade_cell(cell, "F3F4F6")
        _cell_borders(cell, size=4, color="C9CED6")
        cell.paragraphs[0]._p.getparent().remove(cell.paragraphs[0]._p)
        for line in text.strip('\n').split('\n'):
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf = p.paragraph_format
            pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            r = p.add_run(line if line else ' ')
            _style_run(r, rtl=False, size=9, font=CODE_FONT)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def formula(self, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(text)
        _style_run(r, rtl=False, italic=True, size=12, font="Cambria Math")

    def english(self, text, bold=False, size=None, align='justify'):
        p = self.doc.add_paragraph()
        p.alignment = (WD_ALIGN_PARAGRAPH.JUSTIFY if align == 'justify'
                       else WD_ALIGN_PARAGRAPH.CENTER)
        # Explicitly left-to-right paragraph inside the RTL document.
        ppr = p._p.get_or_add_pPr()
        bidi = OxmlElement('w:bidi')
        bidi.set(qn('w:val'), '0')
        _insert_before_any(ppr, bidi, _PPR_AFTER_BIDI)
        r = p.add_run(text)
        _style_run(r, rtl=False, bold=bold, size=size or BODY_PT, font=EN_FONT)
        return p

    # -- captions ----------------------------------------------------------
    def caption(self, label, text):
        if label == FIG_LABEL:
            self.fig_no += 1
            number = self.fig_no
            self.figures.append(f"{label} {number}: {text}")
        else:
            self.tab_no += 1
            number = self.tab_no
            self.tables.append(f"{label} {number}: {text}")
        p = self.doc.add_paragraph(style='Caption')
        _set_bidi(p._p.get_or_add_pPr())
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{label} ")
        _style_run(r, rtl=True)
        add_simple_field(p, f'SEQ {label} \\* ARABIC', str(number), rtl=False)
        r = p.add_run(": ")
        _style_run(r, rtl=True)
        self._add_mixed_text(p, text)
        return p

    # -- tables ------------------------------------------------------------
    def table(self, header, rows, caption, widths=None, font_size=10):
        """RTL data table: caption above (academic convention for tables)."""
        cap = self.caption(TAB_LABEL, caption)
        _keep_with_next(cap)
        table = self.doc.add_table(rows=1, cols=len(header))
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_bidi(table)
        for i, h in enumerate(header):
            cell = table.rows[0].cells[i]
            _shade_cell(cell, HEADER_FILL)
            p = cell.paragraphs[0]
            _set_bidi(p._p.get_or_add_pPr())
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            self._add_mixed_text(p, h, bold=True, size=font_size,
                                 color=RGBColor(0xFF, 0xFF, 0xFF))
        _repeat_header(table.rows[0])
        for ri, row in enumerate(rows):
            cells = table.add_row().cells
            for i, value in enumerate(row):
                if ri % 2 == 1:
                    _shade_cell(cells[i], "F3F6FA")
                p = cells[i].paragraphs[0]
                _set_bidi(p._p.get_or_add_pPr())
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(2)
                self._add_mixed_text(p, str(value), size=font_size)
        if widths:
            for row in table.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return table

    # -- diagrams (native Word shapes built from tables) ---------------------
    def box(self, title, body=None, width_cm=14.0, fill="DCE8F7", border="1F2A44",
            dashed=False, height_lines=0):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_bidi(table)
        cell = table.rows[0].cells[0]
        cell.width = Cm(width_cm)
        _shade_cell(cell, fill)
        _cell_borders(cell, style="dashed" if dashed else "single", size=10, color=border)
        p = cell.paragraphs[0]
        _set_bidi(p._p.get_or_add_pPr())
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(3)
        self._add_mixed_text(p, title, bold=True, size=11, color=NAVY)
        if body:
            for line in body if isinstance(body, (list, tuple)) else [body]:
                bp = cell.add_paragraph()
                _set_bidi(bp._p.get_or_add_pPr())
                bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                bp.paragraph_format.line_spacing = 1.0
                bp.paragraph_format.space_after = Pt(2)
                self._add_mixed_text(bp, line, size=9.5)
        for _ in range(height_lines):
            cell.add_paragraph()
        return table

    def arrow(self, text="▼", note=None):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        _keep_with_next(p)
        r = p.add_run(text)
        _style_run(r, rtl=False, size=12, color=ACCENT)
        if note:
            self._add_mixed_text(p, "  " + note, size=9, color=RGBColor(0x52, 0x51, 0x4E))

    def row_of_boxes(self, items, fill="E3F4EC", font_size=9.5):
        """A horizontal row of boxes separated by arrow columns (reads right-to-left)."""
        cols = len(items) * 2 - 1
        table = self.doc.add_table(rows=1, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _table_bidi(table)
        for i in range(cols):
            cell = table.rows[0].cells[i]
            p = cell.paragraphs[0]
            _set_bidi(p._p.get_or_add_pPr())
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            if i % 2 == 0:
                title, sub = items[i // 2]
                _shade_cell(cell, fill)
                _cell_borders(cell, size=8)
                self._add_mixed_text(p, title, bold=True, size=font_size, color=NAVY)
                if sub:
                    sp = cell.add_paragraph()
                    _set_bidi(sp._p.get_or_add_pPr())
                    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    sp.paragraph_format.line_spacing = 1.0
                    self._add_mixed_text(sp, sub, size=font_size - 1.5)
                cell.width = Cm(2.6)
            else:
                r = p.add_run("◄")
                _style_run(r, rtl=False, size=12, color=ACCENT)
                cell.width = Cm(0.6)
        return table

    def spacer(self, pt=4):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(" ")
        _style_run(r, size=pt)

    def screenshot_placeholder(self, what):
        self.box(f"[ מקום לצילום מסך: {what} ]",
                 "יש להחליף מסגרת זו בצילום מסך של המערכת בעת ההגשה הסופית.",
                 width_cm=15.0, fill="FAFAFA", border="9AA3AF", dashed=True,
                 height_lines=5)

    # -- TOC / lists -------------------------------------------------------
    def field_block(self, instruction, entries, empty_text):
        """A multi-paragraph field whose pre-rendered result lists `entries`.
        Word replaces the result (adding page numbers) when fields are updated."""
        entries = entries or [(1, empty_text)]
        paragraphs = []
        for level, text in entries:
            p = self.doc.add_paragraph()
            _set_bidi(p._p.get_or_add_pPr())
            p.paragraph_format.left_indent = Cm(0.8 * (level - 1))
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            paragraphs.append((p, level, text))
        first, last = paragraphs[0][0], paragraphs[-1][0]
        r = first.add_run()
        _fld_char(r, 'begin', dirty=True)
        r = first.add_run()
        _instr(r, instruction)
        r = first.add_run()
        _fld_char(r, 'separate')
        for p, level, text in paragraphs:
            self._add_mixed_text(p, text, bold=(level == 1))
        r = last.add_run()
        _fld_char(r, 'end')

    def save(self, path):
        self.doc.save(path)


# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------
def title_page(R):
    for _ in range(3):
        R.spacer(14)
    R.para("המכללה האקדמית תל-חי", align='center', bold=True, size=18, color=NAVY)
    R.para("הפקולטה למדעים — החוג למדעי המחשב", align='center', size=14)
    for _ in range(3):
        R.spacer(14)
    R.para("ספר פרויקט גמר", align='center', size=16, color=ACCENT)
    p = R.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("TrendAnalyzer Pro")
    _style_run(r, rtl=False, bold=True, size=32, font=EN_FONT, color=NAVY)
    R.para("מערכת לאיסוף, סינון וניתוח רב-פלטפורמי של מגמות טכנולוגיות "
           "בקהילות מפתחים", align='center', size=15)
    R.english("A Multi-Platform System for Collecting, Filtering and Analysing "
              "Technology Trends in Developer Communities", size=12, align='center')
    for _ in range(4):
        R.spacer(14)
    R.para("מגיש: אלון קרייזברגר (Alon Kreisberger)", align='center', size=14)
    R.para("מנחה: ____________________", align='center', size=14)
    R.spacer(14)
    R.para("תשפ\"ז — 2026", align='center', size=13)


def front_matter(R):
    R.page_break()
    p = R.para("תוכן עניינים", align='center', bold=True, size=18, color=NAVY)
    _keep_with_next(p)
    R.field_block('TOC \\o "1-3" \\h \\z \\u', R.headings,
                  "יש לעדכן שדה זה (F9) כדי להציג את תוכן העניינים.")
    R.page_break()
    R.para("רשימת איורים", align='center', bold=True, size=16, color=NAVY)
    R.field_block(f'TOC \\h \\z \\c "{FIG_LABEL}"', [(2, t) for t in R.figures],
                  "אין איורים.")
    R.spacer(10)
    R.para("רשימת טבלאות", align='center', bold=True, size=16, color=NAVY)
    R.field_block(f'TOC \\h \\z \\c "{TAB_LABEL}"', [(2, t) for t in R.tables],
                  "אין טבלאות.")


def abstract(R):
    R.heading("תקציר", 1, page_break=True)
    R.para(
        "השיח הטכנולוגי העכשווי מבוזר בין קהילות רבות שאינן מדברות זו עם זו: קוד "
        "פתוח נכתב ב-GitHub, דיון תעשייתי מתנהל ב-Hacker News, מדריכים מעשיים "
        "מתפרסמים ב-Dev.to ושיח ציבורי קצר-מועד מתקיים ב-Mastodon. מי שמבקש לזהות "
        "מגמה טכנולוגית אמיתית נדרש לעקוב ידנית אחר כל המקורות, להתמודד עם יחידות "
        "מדידה שאינן ברות השוואה (כוכבים, נקודות, תגובות) ועם כמות רבה של רעש — "
        "דפי חסימה, ספאם, כפילויות ותוכן שאינו קשור לנושא.")
    R.para(
        "בפרויקט זה פותחה מערכת **TrendAnalyzer Pro**, הבנויה בארכיטקטורה תלת-שכבתית: "
        "שכבת איסוף אסינכרונית (Python asyncio ו-httpx) המפעילה במקביל ארבעה אוספים "
        "ייעודיים; שכבת נתונים מבוססת SQLite במצב WAL המאחסנת את הפוסטים בסכמה "
        "אחידה; ושכבת תצוגה אינטראקטיבית ב-Streamlit. בלב המערכת עומד \"שומר הסף\" "
        "(Gatekeeper) — מסנן איכות בן חמישה שלבים (מבנה, רלוונטיות, ייחודיות, שפה "
        "ונושאים) המופעלים מהזול ליקר, כך שמודלי השפה הכבדים רואים רק טקסט ששרד את "
        "הבדיקות הזולות. כל דחייה נרשמת לפי סיבתה, ומשפך הסינון מוצג למשתמש במלואו.")
    R.para(
        "הפוסטים שעברו את הסינון מועשרים בביטויי מפתח באמצעות KeyBERT ובציון סנטימנט "
        "באמצעות VADER, וציוני המעורבות מנורמלים לסקאלה אחידה של 0–100 לכל פלטפורמה "
        "בנפרד. לאחר מכן נבנה גרף דמיון סמנטי בין הפוסטים (שיכוני MiniLM ודמיון "
        "קוסינוס עם ספים מרובדים), הקהילות בגרף מזוהות במקסום מודולריות, ומאותרים "
        "\"גשרים מתהווים\" — זוגות פוסטים מקהילות רחוקות מבנית העוסקים באותו רעיון "
        "תוך שיתוף שכנים מועט, ברוח תאוריית הקשרים החלשים ותאוריית הפערים המבניים. "
        "התוצאות מוצגות בארבע לשוניות: גלגל המערכת האקולוגית, הייפ מול מציאות, "
        "חיבורים, ומסנן שומר הסף, לצד עובד רקע אוטונומי המגלה ועוקב אחר נושאים "
        "חמים מדי שעתיים.")
    R.para("מילות מפתח: זיהוי מגמות, איסוף נתונים אסינכרוני, סינון איכות, KeyBERT, "
           "VADER, ניתוח גרפים, גשרים בין קהילות, Streamlit.", bold=False)

    R.spacer(8)
    R.english("Abstract", bold=True, size=14, align='center')
    R.english(
        "Technology discourse is fragmented across communities that rarely overlap: "
        "open-source code lives on GitHub, industry debate on Hacker News, practitioner "
        "write-ups on Dev.to and short-form public discussion on Mastodon. Identifying a "
        "genuine trend therefore requires manually following every source, reconciling "
        "incomparable engagement units (stars, points, reactions) and discarding a large "
        "amount of noise such as bot walls, spam, duplicates and off-topic content.")
    R.english(
        "This project presents TrendAnalyzer Pro, a system built on a three-tier "
        "architecture: an asynchronous collection tier (Python asyncio and httpx) that runs "
        "four platform collectors concurrently; a persistence tier based on SQLite in WAL "
        "mode that stores posts in a unified schema; and an interactive presentation tier "
        "built with Streamlit. At its core is the Gatekeeper, a five-stage quality filter "
        "(structure, relevance, uniqueness, language and themes) ordered cheapest-first so "
        "that expensive language models only see text that survived the inexpensive checks. "
        "Every rejection is tallied by reason, and the complete funnel is exposed to the user.")
    R.english(
        "Accepted posts are enriched with KeyBERT key phrases and VADER sentiment, and raw "
        "engagement is normalised to a per-platform 0-100 trend score. A semantic similarity "
        "graph is then built from MiniLM sentence embeddings with stratified cosine "
        "thresholds; communities are detected by modularity maximisation, and emerging "
        "bridges, i.e. pairs of posts from structurally distant communities that discuss "
        "the same idea while sharing few neighbours, are ranked in the spirit of weak-tie "
        "and structural-hole theory. Results are presented in four dashboard views "
        "(Ecosystem Wheel, Hype vs. Reality, Connections and the Gatekeeper's Filter), "
        "complemented by an autonomous background worker that discovers and tracks "
        "trending topics every two hours.")
    R.english("Keywords: trend detection, asynchronous data collection, data quality "
              "filtering, KeyBERT, VADER, graph analysis, community bridges, Streamlit.")


def chapter1(R):
    R.heading("פרק 1: מבוא", 1, page_break=True)
    R.heading("1.1 רקע", 2)
    R.para(
        "בעשור האחרון הפך תחום פיתוח התוכנה לאחד התחומים הדינמיים ביותר בכלכלה. "
        "טכנולוגיות, ספריות ומתודולוגיות חדשות צצות בקצב של שבועות, וחלקן נעלמות "
        "באותה מהירות. החלטות של ארגונים — באילו כלים להשקיע, אילו מיומנויות לפתח "
        "ואילו מוצרים לבנות — תלויות ביכולת להבחין בין מגמה בעלת ממשות לבין \"באזז\" "
        "חולף. ספרות זיהוי המגמות עוסקת בשאלה זו שנים רבות, החל במודלים של \"פרצי "
        "פעילות\" בזרמי טקסט (Kleinberg, 2003) ועד מערכות לזיהוי מגמות בזמן אמת "
        "ברשתות חברתיות (Mathioudakis & Koudas, 2010).")
    R.para(
        "דא עקא, המידע הרלוונטי למפתחים אינו מרוכז ברשת חברתית אחת. כל קהילה מייצגת "
        "\"משלב\" (register) שונה של השיח: ב-GitHub מתבטא השיח בקוד עובד ובמאגרים "
        "מתוחזקים; ב-Hacker News — בדיון תעשייתי סביב מאמרים מקושרים; ב-Dev.to — "
        "במדריכים ובחוויות מעשיות של מתרגלים; וב-Mastodon — בשיח ציבורי קצר. מגמה "
        "אמיתית נוטה להופיע בכמה מן המשלבים הללו בו-זמנית, ואילו רעש נוטה להישאר "
        "בקהילה אחת.")

    R.heading("1.2 מוטיבציה", 2)
    R.para(
        "המוטיבציה לפרויקט נובעת משלושה קשיים מעשיים שעמם מתמודד כל מי שמנסה לעקוב "
        "אחר מגמות טכנולוגיות:")
    R.bullets([
        "**פיזור:** אין ממשק אחד המאחד את המקורות, ולכל פלטפורמה ממשק תכנות (API) "
        "שונה, תחביר חיפוש שונה ומגבלות קצב שונות.",
        "**חוסר השוואתיות:** 200 כוכבים ב-GitHub ו-200 נקודות ב-Hacker News אינם "
        "מבטאים אותה רמת עניין; יש צורך בנרמול סטטיסטי לפי פלטפורמה.",
        "**רעש:** חלק ניכר מתוצאות החיפוש הגולמיות הוא דפי CAPTCHA, חומות תשלום, "
        "פוסטים של האשטגים בלבד, כפילויות, תוכן שאינו באנגלית או תוכן שאינו קשור "
        "לנושא כלל. ניתוח ללא סינון מוביל למסקנות שגויות.",
    ])
    R.para(
        "מעבר לכך, ערך מוסף משמעותי טמון דווקא בקשרים שבין הקהילות: כאשר מאגר קוד "
        "ב-GitHub ודיון ב-Mastodon עוסקים באותו רעיון אף שאין ביניהם קשר ישיר, ייתכן "
        "שמדובר במגמה בשלב התהוות. זיהוי \"גשרים\" כאלה הוא אחת מתרומותיו העיקריות "
        "של הפרויקט.")

    R.heading("1.3 הגדרת הבעיה", 2)
    R.para(
        "בהינתן נושא טכנולוגי שמזין המשתמש (למשל \"rust\" או \"AI agents\"), יש "
        "לאסוף ממספר פלטפורמות הטרוגניות פוסטים עדכניים ורלוונטיים, לסנן מהם רעש "
        "באופן שקוף ומדיד, לנרמל את מדדי המעורבות לסקאלה אחידה, ולהפיק תובנות "
        "ניתנות להסבר: היכן מתרכזת תשומת הלב, האם השיח מגובה בקוד ממשי, ואילו "
        "רעיונות חוצים את גבולות הקהילות. כל זאת בזמן תגובה אינטראקטיבי (פחות "
        "משלוש דקות לסריקה מלאה) ובמגבלות הקצב של ממשקי התכנות הציבוריים.")
    R.para("מן ההגדרה נגזרות הדרישות הבאות:")
    R.bullets([
        "איסוף מקבילי מארבע פלטפורמות לפחות, עם בידוד תקלות: כשל של פלטפורמה אחת "
        "אינו מפיל את הסריקה כולה.",
        "סינון איכות דטרמיניסטי, שאת ספיו אי-אפשר להחליש מבחוץ, ושכל דחייה בו "
        "נרשמת עם סיבה.",
        "ייצוג מאוזן של הקהילות: פלטפורמה בעלת מעורבות גבוהה באופן שיטתי לא תדחק "
        "את האחרות.",
        "ניתוח סמנטי וגרפי הניתן להסבר, ללא תלות בשירות חיצוני בתשלום.",
        "ממשק משתמש אינטראקטיבי המציג גם את התוצאות וגם את תהליך הסינון שהוביל "
        "אליהן.",
    ], numbered=True)

    R.heading("1.4 מטרות הפרויקט ותרומתו", 2)
    R.para(
        "מטרת הפרויקט היא לתכנן ולממש מערכת מקצה לקצה העונה על הדרישות לעיל. "
        "תרומות הפרויקט העיקריות הן: (א) ארכיטקטורת איסוף אסינכרונית עם דגימת יתר "
        "וגלי עיבוד אדפטיביים; (ב) מנגנון Gatekeeper חמישה-שלבי עם חשבונאות משפך "
        "מאוזנת; (ג) שיטת נרמול לוגיסטית לכל פלטפורמה; (ד) אלגוריתם לזיהוי גשרים "
        "מתהווים בגרף דמיון סמנטי המשלב דמיון, שיבוץ-חלש ומרחק משלב; ו-(ה) לוח מחוונים "
        "המציג את כל אלה בצורה שקופה.")

    R.heading("1.5 מבנה הספר", 2)
    R.para(
        "פרק 2 מתאר את ארכיטקטורת המערכת התלת-שכבתית. פרק 3 מפרט את מנגנון "
        "Gatekeeper. פרק 4 עוסק בעיבוד השפה הטבעית ובניתוח הגרפים. פרק 5 מציג את "
        "ממשק המשתמש, ופרק 6 מסכם, דן באתגרים ומציע כיווני המשך.")


def chapter2(R):
    R.heading("פרק 2: ארכיטקטורת המערכת", 1, page_break=True)
    R.heading("2.1 סקירה כללית", 2)
    R.para(
        "המערכת בנויה בארכיטקטורה תלת-שכבתית קלאסית (Fowler, 2002), שבה כל שכבה "
        "תלויה רק בשכבה שמתחתיה: שכבת האיסוף והלוגיקה (collectors/ ו-engine/), "
        "שכבת הנתונים (database/) ושכבת התצוגה (ui/). קובץ התצורה config.py משמש "
        "\"מקור אמת יחיד\" לכל הקבועים הניתנים לכוונון, והוא נקי מתופעות לוואי, כך "
        "שכל שכבה יכולה לייבא אותו בבטחה. המערכת מופעלת באמצעות main.py, המאפס את "
        "בסיס הנתונים ומפעיל את Streamlit כתהליך-בן; לוח המחוונים מסרב לרוץ ללא "
        "דגל ההפעלה TREND_APP_LAUNCHED, כך שאיפוס הנתונים קודם תמיד להצגתם.")

    R.box("שכבת התצוגה — Presentation Tier (ui/)",
          ["Streamlit: app.py, views/ (home, search, 4 tabs), data.py, scan.py",
           "גרף סמנטי וזיהוי גשרים: graph_analyzer.py, graph_bridges.py"],
          fill=BAND_FILLS[0])
    R.arrow("▲▼", "קריאה בלבד (WAL) / הפעלת סריקה")
    R.box("שכבת הלוגיקה והאיסוף — Ingestion & Engine Tier",
          ["engine/: validation → planning → cycle → scoring · sentiment · signals",
           "collectors/: GitHub · Hacker News · Dev.to · Mastodon + Gatekeeper (base.py)",
           "scheduler.py: עובד רקע אוטונומי (APScheduler)"],
          fill=BAND_FILLS[1])
    R.arrow("▲▼", "כתיבה טרנזקציונית")
    R.box("שכבת הנתונים — Persistence Tier (database/)",
          ["SQLite (WAL): unified_posts · scan_volume · live_trends · live_posts",
           "TrendManager (manager.py), schema.py, live.py"],
          fill=BAND_FILLS[2])
    R.caption(FIG_LABEL, "הארכיטקטורה התלת-שכבתית של TrendAnalyzer Pro")

    R.heading("2.2 מחזור סריקה מקצה לקצה", 2)
    R.para(
        "סריקה אינטראקטיבית מתחילה כאשר המשתמש מזין נושא בשורת החיפוש. הפונקציה "
        "engine.scan() מאמתת את הנושא, נועלת את מנעול הסריקה הגלובלי (_SCAN_LOCK) "
        "ומריצה את run_cycle() בתוך לולאת asyncio. השלבים מדווחים ללוח המחוונים "
        "באמצעות פונקציית התקדמות (callback), כך שהמנוע עצמו אינו תלוי ב-Streamlit "
        "וניתן להפעילו גם מסקריפטים ומבדיקות.")
    R.row_of_boxes([
        ("אימות", "validate_topic"),
        ("תכנון", "TopicProfile"),
        ("איסוף מקבילי", "4 collectors"),
        ("Gatekeeper", "5 stages"),
        ("נרמול וגיזום", "0–100, Top-K"),
        ("אחסון", "SQLite"),
    ])
    R.caption(FIG_LABEL, "שלבי מחזור סריקה אינטראקטיבית (מימין לשמאל)")
    R.bullets([
        "**אימות קלט** (engine/validation.py): הנושא מנורמל (כיווץ רווחים, עד 80 "
        "תווים) ונבדק מול אלפבית אנגלי מותר הכולל סימנים הנפוצים בשמות טכנולוגיות "
        "(c++, node.js, ci/cd). האימות מתבצע פעמיים — בממשק ובתוך המנוע — כך שאף "
        "קורא אינו יכול לעקוף אותו.",
        "**תכנון** (engine/planning.py, engine/topic_profile.py): DynamicTopicManager "
        "הופך את מילות הנושא לתוכנית חיפוש דטרמיניסטית — שאילתות לכל פלטפורמה ואוצר "
        "מילים דו-שכבתי לרלוונטיות. כל ערך עובר סניטציה (תבנית slug ל-GitHub, תגיות "
        "אלפא-נומריות ל-Mastodon ו-Dev.to, עד 8 שאילתות, 6 תגיות ו-40 מילות מפתח), "
        "כך שהנושא אינו יכול להזריק תחביר שאילתה.",
        "**איסוף מקבילי** (engine/cycle.py): ארבעת האוספים רצים ב-asyncio.gather "
        "על לקוח httpx.AsyncClient משותף (איגום חיבורים), וכל אחד עטוף ב-"
        "asyncio.wait_for עם תקרה של 150 שניות. משך הסריקה שווה אפוא למשך הפלטפורמה "
        "האיטית ביותר, ולא לסכום כולן.",
        "**מסנני רעש:** פלטפורמה שהניבה פחות מ-5 פוסטים נשמטת, וסריקה שסך הפוסטים "
        "בה נמוך מ-15 או שהגיעה לפחות משלוש פלטפורמות — מבוטלת לפני כל כתיבה לבסיס "
        "הנתונים, מתוך ההנחה שמגמה היא בהגדרתה תופעה חוצת-קהילות.",
        "**ניקוד ואחסון** (engine/scoring.py): נושא חדש מחליף את הקורפוס באופן "
        "אטומי; ציוני המעורבות מנורמלים לכל פלטפורמה, ומתבצע גיזום מרובד.",
    ], numbered=True)

    R.heading("2.3 שכבת האיסוף", 2)
    R.para(
        "כל אוסף יורש מן המחלקה המופשטת BaseCollector (collectors/base.py) "
        "ומממש מתודה אחת בלבד — collect(client) — המגדירה כיצד לתשאל את הפלטפורמה, "
        "כיצד לבצע \"שליפה עמוקה\" (deep fetch) של התוכן המלא, וכיצד למפות את הרשומה "
        "לסכמת הפוסט האחידה. כל מה שקובע אם פוסט יישמר, וכל החשבונאות המסבירה את "
        "ההחלטה, ממומשים פעם אחת במחלקת הבסיס, המורכבת משני Mixins: RelevanceMixin "
        "(מסנן האיכות) ו-FunnelMixin (מאגר המועמדים, גלי העיבוד ודוח הנפח). זהו "
        "יישום של תבנית Template Method.")
    R.table(
        ["פלטפורמה", "ממשק ושאילתה", "שליפה עמוקה", "מקביליות", "אות מעורבות"],
        [
            ["GitHub", "Search API — topic slugs, דחיפה ב-14 הימים האחרונים, "
             "רצפת כוכבים 200/100/50 לפי דירוג הנושא",
             "README דרך raw.githubusercontent (CDN) ואז API; הסרת Markdown",
             "Semaphore(5), שמירת מכסה 15", "כוכבים"],
            ["Hacker News", "Algolia Search API — טקסט חופשי, 7 ימים, רצפת נקודות 5 "
             "(2 בסריקה מצומצמת)",
             "סריקת הקישור החיצוני עם BeautifulSoup, עד 2MB לדף",
             "Semaphore(8)", "נקודות"],
            ["Dev.to", "Forem API — תגית אחת לבקשה, המאמרים המובילים של השבוע",
             "גוף המאמר המלא ב-Markdown לכל מועמד", "Semaphore(2)", "תגובות (reactions)"],
            ["Mastodon", "Tag timeline ב-mastodon.social, עד 8 עמודים של 40 "
             "(max_id)",
             "לא נדרשת — הטוט מגיע שלם; סינון מטא-דאטה מקדים",
             "סדרתי לפי עמוד", "מועדפים והדהודים"],
        ],
        "השוואת ארבעת אוספי הפלטפורמות", widths=[2.3, 4.6, 4.2, 2.4, 2.2], font_size=9)

    R.heading("2.3.1 דגימת יתר וגלי עיבוד אדפטיביים", 3)
    R.para(
        "מכיוון ש-Gatekeeper דוחה חלק ניכר מן התוצאות הגולמיות, כל אוסף מבקש מאגר "
        "מועמדים גדול פי ארבעה מן היעד (CANDIDATE_OVERSAMPLE = 4; 100 מועמדים ליעד "
        "של 25 פוסטים). הפונקציה merge_candidates ממזגת את תוצאות השאילתות, מסירה "
        "כפילויות לפי מזהה ומדרגת לפי מעורבות. לאחר מכן collect_until_target "
        "צורכת את המאגר \"בעצלות\" בגלים מקביליים: הגל הראשון בגודל 30, וגודל כל "
        "גל נוסף מחושב משיעור הקבלה שנצפה עד כה, עם רצפה של 0.2 כדי שגל שלא קיבל "
        "דבר לא יבקש את כל המאגר. הסריקה נעצרת מיד כשהיעד הושג, כך שפלטפורמה בעלת "
        "שיעור קבלה גבוה \"משלמת\" על גל אחד בלבד.")
    R.code('''
results = await asyncio.gather(
    *[process(item) for item in batch], return_exceptions=True)
...
rate = max(len(accepted) / cursor, 0.2)
wave_size = min(len(candidates) - cursor,
                int((target - len(accepted)) / rate) + 2)''')
    R.para(
        "השימוש ב-return_exceptions=True מממש בידוד תקלות ברמת הפריט: חריגה בעיבוד "
        "פוסט אחד נרשמת כ-processing_error ואינה מבטלת את שאר הגל.")

    R.heading("2.3.2 בקרת מקביליות ומגבלות קצב", 3)
    R.para(
        "לכל מארח הוגדר asyncio.Semaphore משלו: חמש בקשות מקביליות ל-GitHub, שתיים "
        "ל-Dev.to (שאוכף מגבלה מחמירה ללקוח) ושמונה לדפים חיצוניים המקושרים "
        "מ-Hacker News. אוסף GitHub בודק את מכסת ה-API (60 בקשות בשעה ללא אסימון, "
        "5,000 עם GITHUB_TOKEN) ומפסיק להוציא מכסה על שליפת README דרך ה-API כאשר "
        "נותרו פחות מ-15 בקשות. כל בקשת HTTP מוגבלת ל-15 שניות.")

    R.heading("2.3.3 חשבונאות המשפך והערכת נפח", 3)
    R.para(
        "מכיוון שכל אוסף נעצר ביעד, מספר הפוסטים שנשמרו אינו מעיד על הנפח הקיים "
        "בפלטפורמה. לכן כל אוסף מפיק volume_report המעריך את מספר הפריטים בפלטפורמה "
        "שתואמים לחיפוש ועשויים לעבור את המסנן:")
    R.formula("estimate = searched_total × overlap × acceptance")
    R.para(
        "כאשר searched_total הוא סכום הספירות המדווחות על-ידי הפלטפורמה, overlap הוא "
        "יחס הייחודיים לשורות שנשלפו (תיקון לחפיפה בין שאילתות קרובות), ו-acceptance "
        "הוא שיעור הקבלה של המסנן. המשפך מקיים תמיד את האיזון gated = accepted + "
        "Σ reject_reasons; מועמדים שנכשלו לפני שהמסנן יכול היה לשפוט אותם (קישור מת, "
        "רשומה פגומה) מחושבים כשארית ומדווחים כ-unreachable.")

    R.heading("2.4 נרמול ציוני מגמה", 2)
    R.para(
        "יחידות המעורבות שונות בין הפלטפורמות, ולכן הנרמול מתבצע לכל פלטפורמה "
        "בנפרד (recalculate_platform_stats). הלוגריתם מרסן את הזנב הכבד של התפלגות "
        "המעורבות, ציון התקן (Z-score) ממקם כל פוסט ביחס לפלטפורמה שלו, והפונקציה "
        "הלוגיסטית חוסמת את התוצאה לטווח 0–100:")
    R.formula("trend_score = 100 · σ( (log10(raw + 1) − μ_p) / s_p + 0.5 )")
    R.para(
        "ההיסט 0.5 ממקם פוסט ממוצע בכ-62 נקודות ולא בדיוק ב-50, ופיזור אפס (כל "
        "הציונים זהים) מוחלף ב-1 כדי למנוע חלוקה באפס. העדכון כולו מבוצע בטרנזקציה "
        "אחת. לאחר הנרמול מבוצע גיזום מרובד (prune_database): רק 25 הפוסטים "
        "המובילים בכל פלטפורמה נשמרים, בשאילתה יחידה המשתמשת ב-ROW_NUMBER() OVER "
        "(PARTITION BY source_platform). הגיזום בתוך כל פלטפורמה, ולא באופן גלובלי, "
        "מבטיח ייצוג שווה לכל הקהילות.")

    R.heading("2.5 שכבת הנתונים", 2)
    R.para(
        "שכבת הנתונים ממומשת במחלקה TrendManager (database/manager.py), המשמשת "
        "Data Access Object. כל מתודה ציבורית פותחת חיבור קצר-מועד משלה, כך שאובייקט "
        "יחיד ניתן לשיתוף בטוח בין תהליכון הסריקה לתהליכוני הרינדור של Streamlit. "
        "בסיס הנתונים פועל במצב Write-Ahead Logging, שבו קוראים (לוח המחוונים) "
        "אינם נחסמים על-ידי כותב (סריקה), ונקבע busy timeout של 30 שניות. האילוץ "
        "UNIQUE(source_platform, external_id) הופך את INSERT OR IGNORE לאידמפוטנטי "
        "בין סריקות חוזרות. הסכמה מנוהלת ב-schema.py, כולל הגירות קדימה (ALTER "
        "TABLE ADD COLUMN) לבסיסי נתונים שנוצרו בגרסאות קודמות.")
    R.table(
        ["טבלה", "תוכן", "עמודות עיקריות"],
        [
            ["unified_posts", "שורה לכל פוסט שהתקבל, בסכמה בלתי תלויה בפלטפורמה",
             "source_platform, external_id, title, content, raw_score, trend_score, "
             "published_at, keywords (JSON)"],
            ["scan_volume", "שורה לכל פלטפורמה המתארת את הסריקה האחרונה לפני התקרה",
             "searched_total, overlap, acceptance_rate, estimated_matches, gated, "
             "accepted, reject_reasons (JSON)"],
            ["live_trends", "תוצאות עובד הרקע: שורה לכל פלטפורמה לכל נושא",
             "topic, posts, peak_score, avg_score, matches, sentiment, keywords"],
            ["live_posts", "כל פוסט שנשמר בסריקת רקע, לאחר נרמול",
             "topic, source_platform, title, url, trend_score"],
        ],
        "טבלאות בסיס הנתונים ותפקידן", widths=[2.8, 6.0, 7.0], font_size=9.5)
    R.para(
        "שיכוני הטקסט (embeddings) אינם נשמרים בבסיס הנתונים במכוון: נשמרת רק רשימת "
        "מילות מפתח קומפקטית בפורמט JSON, והגרף הסמנטי מחשב את השיכונים מחדש לפי "
        "דרישה. מדיניות שמירת נתונים (TTL) של שבעה ימים נאכפת בתחילת כל מחזור רקע.")

    R.heading("2.6 עובד הרקע האוטונומי", 2)
    R.para(
        "לצד הסריקה האינטראקטיבית פועל עובד רקע (scheduler.py) המבוסס על "
        "APScheduler ומריץ מחזור כל שעתיים. המחזור כולל שלושה שלבים:")
    R.box("1. פינוי — Eviction", "cleanup_old_records(): מחיקת רשומות ישנות מ-7 ימים",
          width_cm=12, fill=BAND_FILLS[4])
    R.arrow()
    R.box("2. גילוי — Discovery",
          ["40 הכתבות המובילות ב-Hacker News ו-60 מאמרים ב-Dev.to",
           "תגיות המופיעות ב-2 מאמרים לפחות, לאחר רשימת חסימה של תגיות-מטא",
           "KeyBERT מדרג את התגיות לפי מרכזיותן בכותרות → 5 נושאים"],
          width_cm=12, fill=BAND_FILLS[0])
    R.arrow()
    R.box("3. מעקב — Tracking",
          ["engine.scan_background(): אותו צינור, 5 פוסטים לפלטפורמה",
           "נרמול בבסיס נתונים זמני; שמירה ל-live_trends / live_posts"],
          width_cm=12, fill=BAND_FILLS[1])
    R.caption(FIG_LABEL, "מחזור עובד הרקע האוטונומי")
    R.para(
        "סריקת הרקע רצה מול בסיס נתונים זמני (tempfile), כך שציוני ה-Z מחושבים על "
        "הנושא בלבד והקורפוס שהמשתמש חוקר לעולם אינו נפגע. לעובד הרקע מנעול נפרד "
        "(_BACKGROUND_LOCK), ולכן הוא יכול לרוץ במקביל לסריקה אינטראקטיבית. סריקות "
        "רקע משתמשות בספי רעש מופחתים (פוסט אחד לפלטפורמה, שישה בסך הכול), כדי "
        "שקהילות בעלות נפח נמוך עדיין יופיעו בדף הבית. בעת הפעלה ראשונה (Cold Start) "
        "מוצג מסך טעינה עד שנאספו חמישה נושאים, וכל מחזור מתזמן השלמה עבור נושאים "
        "שטרם נוסו.")

    R.heading("2.7 מקביליות ובטיחות תהליכונים", 2)
    R.para(
        "המערכת משלבת שני מודלי מקביליות: asyncio בתוך סריקה, ותהליכונים (threads) "
        "בין הסריקה האינטראקטיבית, עובד הרקע וסשנים של Streamlit. מכאן נגזרו כמה "
        "החלטות תכנון:")
    R.bullets([
        "מודל KeyBERT ומקודד המשפטים נטענים כ-Singleton ברמת התהליך, וכל קריאה "
        "אליהם מוגנת במנעול (_KW_LOCK, _ENCODER_LOCK), משום שהמודל והטוקנייזר "
        "המהיר אינם תומכים בקריאות מקביליות.",
        "ההסקה מוצמדת ל-CPU (EMBED_DEVICE = \"cpu\"), משום ש-backend ה-MPS של Apple "
        "אינו בטוח לשימוש מכמה תהליכונים, וקידוד מקבילי עלול להשחית את זיכרון ה-GPU.",
        "Streamlit מריץ סקריפטים בתהליכון ללא לולאת אירועים; לכן _run_coroutine "
        "משתמש ב-asyncio.run, ובמקרה שקיימת כבר לולאה — מריץ את הקורוטינה בלולאה "
        "פרטית על תהליכון עובד.",
        "ביטול שיתופי: run_cycle בודק דגל threading.Event כל חצי שנייה, ומבטל את "
        "האוספים לפני שנכתב דבר לבסיס הנתונים.",
    ])


def chapter3(R):
    R.heading("פרק 3: מנגנון שומר הסף (The Quality Gatekeeper)", 1, page_break=True)
    R.heading("3.1 עקרונות התכנון", 2)
    R.para(
        "Gatekeeper הוא הרכיב המכריע אם פוסט מנורמל יישמר. הוא ממומש במתודה "
        "is_quality_content של RelevanceMixin (collectors/relevance.py), ומשותף לכל "
        "ארבעת האוספים כך שההחלטה זהה בכל פלטפורמה. התכנון נשען על ארבעה עקרונות:")
    R.bullets([
        "**מהזול ליקר:** הבדיקות המבניות הזולות רצות ראשונות, ורק טקסט ששרד אותן "
        "מגיע לזיהוי השפה ולאחר מכן ל-KeyBERT — השלב היקר ביותר מבחינת CPU.",
        "**סיבה אחת לכל דחייה:** השלב הראשון שנכשל רושם את סיבתו ב-reject_reasons "
        "ועוצר את ההערכה, כך שכל פוסט שנדחה נספר בדיוק פעם אחת.",
        "**קשיחות בלתי ניתנת להחלשה:** תוכנית החיפוש רשאית לשנות מה נחשב רלוונטי "
        "(אוצר המילים), אך לעולם לא את מידת הקפדנות; כל הספים המבניים קבועים "
        "ב-config.py. שכבת מילים ריקה מהודרת לביטוי NEVER_MATCH, כך שנושא ללא "
        "אוצר מילים שמיש אינו מקבל דבר — ולא חוזר בשקט לנושא ברירת המחדל.",
        "**שקיפות:** כל ספירות הדחייה זורמות ל-volume_report ומוצגות בלשונית "
        "Gatekeeper's Filter.",
    ])
    R.para(
        "לפני השלבים עצמם מבוצע ניקוי (clean_text): פענוח ישויות HTML, החלפת תגיות "
        "ברווחים, הסרת כתובות URL וכיווץ רווחים, כך שהמסנן מודד את מה שהקורא רואה. "
        "תוכן Markdown (README, מאמרי Dev.to) עובר לפני כן clean_markdown, המסיר "
        "frontmatter, בלוקי קוד, תגי badge, תמונות, טבלאות וכותרות.")

    widths = (15.5, 13.5, 11.5, 9.5, 7.5)
    stages = [
        ("שלב 1 — מבנה (Structure)",
         "junk_page · too_short · too_few_words · low_alpha_ratio · hashtag_soup · repetitive"),
        ("שלב 2 — רלוונטיות (Relevance)", "not_topic_relevant"),
        ("שלב 3 — ייחודיות (Uniqueness)", "duplicate — SHA-256"),
        ("שלב 4 — שפה (Language)", "not_english · lang_undetectable"),
        ("שלב 5 — נושאים (Themes)", "no_keywords — KeyBERT"),
    ]
    for i, (title, body) in enumerate(stages):
        R.box(title, body, width_cm=widths[i], fill=BAND_FILLS[i % len(BAND_FILLS)])
        R.arrow("▼" if i < len(stages) - 1 else "▼", None if i < len(stages) - 1
                else "פוסט מאושר: keywords + published_at (ISO 8601)")
    R.caption(FIG_LABEL, "משפך חמשת השלבים של Gatekeeper וקודי הדחייה בכל שלב")

    R.heading("3.2 שלב 1: בדיקות מבנה", 2)
    R.para(
        "השלב המבני (_structural_reject) מופעל על גוף הפוסט בלבד, ללא הכותרת, משום "
        "שבפלטפורמות מסוימות (Mastodon) הכותרת נגזרת מן הגוף והטקסט היה נספר פעמיים. "
        "כל בדיקה מכוונת למחלקה אחרת של קלט דל-ערך:")
    R.table(
        ["קוד דחייה", "תנאי", "סף (config.py)", "מטרה"],
        [
            ["junk_page", "התאמה לתבנית חסימה ב-400 התווים הראשונים",
             "JUNK_PATTERNS (10 תבניות)", "CAPTCHA, חומות תשלום ודפי שגיאה המוחזרים "
             "עם HTTP 200"],
            ["too_short", "אורך הגוף קטן מהסף", "MIN_CONTENT_CHARS = 180",
             "תוכן שאין בו די לניתוח (GitHub רשאי להנמיך)"],
            ["too_few_words", "מספר המילים קטן מהסף (מותאם לסף האורך, לפחות 8)",
             "MIN_WORD_COUNT = 30", "קטעים קצרים וכותרות בלבד"],
            ["low_alpha_ratio", "שיעור האותיות והרווחים קטן מהסף",
             "MIN_ALPHA_RATIO = 0.55", "שאריות בינאריות, base64 ו-markup"],
            ["hashtag_soup", "שיעור האסימונים שמתחילים ב-# גבוה מהסף",
             "MAX_HASHTAG_RATIO = 0.25", "פוסטים של האשטגים בלבד"],
            ["repetitive", "יחס המילים הייחודיות ב-200 האסימונים הראשונים (מ-20 "
             "מילים ומעלה) קטן מהסף", "MIN_UNIQUE_TOKEN_RATIO = 0.35",
             "ספאם חזרתי"],
        ],
        "בדיקות השלב המבני וספיהן", widths=[2.8, 5.4, 3.8, 4.2], font_size=9)
    R.para(
        "הבחירה בחלון קבוע של 200 אסימונים לבדיקת החזרתיות נובעת מחוק Heaps (Heaps, "
        "1978): עושר אוצר המילים יורד באופן טבעי ככל שהמסמך מתארך, ולכן מדידה על "
        "המסמך כולו הייתה מענישה טקסטים ארוכים ולגיטימיים. ספאם חזרתי אמיתי ניכר "
        "כבר במאות המילים הראשונות.")

    R.heading("3.3 שלב 2: רלוונטיות לנושא", 2)
    R.para(
        "בדיקת הרלוונטיות (is_topic_relevant) מבוססת על אוצר מילים דו-שכבתי "
        "והתאמה בגבולות מילה (\\b). כל רשימת מונחים מהודרת לביטוי רגולרי יחיד "
        "שאינו רגיש לרישיות, ממוין מהארוך לקצר כך שביטוי רב-מילי גובר על התחילית שלו, "
        "והטקסט נסרק פעם אחת בלבד במקום פעם לכל מונח. עוגני גבול המילה חיוניים: "
        "ללא הם, מונחים קצרים כגון ai או ml היו מותאמים בתוך available או html.")
    R.bullets([
        "**מונחים חזקים** (Strong) — חד-משמעיים; הופעה אחת בכל מקום בטקסט מספיקה. "
        "לדוגמה, בנושא ברירת המחדל: llm, pytorch, fine-tuning, huggingface.",
        "**מונחים חלשים** (Weak) — רב-משמעיים; מתקבלים רק בכותרת (שהיא קצרה ומכוונת), "
        "או כאשר לפחות שני מונחים חלשים שונים מופיעים יחד בטקסט ומאששים זה את זה. "
        "חזרות על אותו מונח חלש אינן נחשבות לאישוש.",
    ])

    R.heading("3.4 שלב 3: ייחודיות", 2)
    R.para(
        "תוכן זהה עשוי להופיע תחת מזהים חיצוניים שונים (למשל מאמר שפורסם מחדש). לכן "
        "מחושבת טביעת אצבע SHA-256 של גוף הפוסט לאחר המרה לאותיות קטנות והסרת כל "
        "התווים שאינם תווי מילה, כך שגם עותקים שעוצבו מחדש באופן טריוויאלי מזוהים. "
        "טביעת האצבע נרשמת בקבוצה _seen_hashes רק עם קבלת הפוסט, כדי שפוסט שנדחה "
        "בשלב מאוחר יותר לא יחסום עותק תקין של אותו טקסט. האילוץ UNIQUE בבסיס "
        "הנתונים משלים את הבדיקה ברמת המזהה החיצוני.")
    R.code('''
content_hash = hashlib.sha256(
    re.sub(r'\\W+', '', post['content'].lower()).encode('utf-8')
).hexdigest()''')

    R.heading("3.5 שלב 4: זיהוי שפה", 2)
    R.para(
        "כל המקורות והתאמת הרלוונטיות מבוססים על אנגלית, ולכן פוסטים בשפות אחרות "
        "נדחים. הזיהוי מבוצע בספריית langdetect (Shuyo, 2010) על 500 התווים "
        "הראשונים של הכותרת והגוף — כמות המספיקה להערכה בטוחה. נדרשת שפה עליונה "
        "\"en\" בהסתברות של 0.85 לפחות (LANG_MIN_PROB). דגימות קצרות מ-25 תווים "
        "מדולגות, משום שהמזהה אינו אמין עליהן, וחריגה של הספרייה נרשמת "
        "כ-lang_undetectable.")

    R.heading("3.6 שלב 5: חילוץ נושאים", 2)
    R.para(
        "השלב האחרון מפעיל את KeyBERT (Grootendorst, 2020) לחילוץ ביטויי מפתח. "
        "פוסט שהמודל אינו מוצא בו אף ביטוי בביטחון מספק נחשב כחסר תוכן בר-ניתוח "
        "ונדחה כ-no_keywords. הפרמטרים נבחרו כך שהביטויים יהיו מגוונים ולא וריאציות "
        "של רעיון אחד:")
    R.code('''
kw_model.extract_keywords(
    text, keyphrase_ngram_range=(1, 2), stop_words='english',
    top_n=8, use_mmr=True, diversity=0.6)
# keep phrases with score > 0.35, at most five''')
    R.para(
        "עם הקבלה, הפוסט מועשר במקום בשדה keywords ובחותמת זמן published_at "
        "מנורמלת ל-ISO 8601 ב-UTC. כשל של המודל דוחה את הפוסט הבודד ולעולם לא את "
        "הסריקה כולה.")

    R.heading("3.7 סינון מקדים ברמת המקור", 2)
    R.para(
        "בנוסף לחמשת השלבים המשותפים, אוסף Mastodon מפעיל מסננים זולים על "
        "המטא-דאטה לפני ש-Gatekeeper רואה את הטקסט: הדהודים (reposts), פוסטים עם "
        "אזהרת תוכן (sensitive), שפה מוצהרת שאינה אנגלית, תגובות אישיות (reply) "
        "ופוסטים המורכבים מהאשטגים בלבד. גם דחיות אלה נרשמות בקוד סיבה משלהן, כך "
        "שהמשפך נשאר מאוזן.")

    R.heading("3.8 קטגוריות דחייה ותצוגתן", 2)
    R.para(
        "לצורך הצגה למשתמש, קודי הדחייה הגולמיים מקובצים לקטגוריות קריאות "
        "(REJECT_CATEGORIES בקובץ ui/views/gatekeeper_tab.py):")
    R.table(
        ["קטגוריה", "קודי דחייה", "הסבר"],
        [
            ["Off-topic", "not_topic_relevant", "לא נמצאו מונחים חזקים ומעט מדי מונחים חלשים"],
            ["Too short or thin", "too_short, too_few_words", "מתחת לסף התווים או המילים"],
            ["Wrong language", "not_english, lang_undetectable", "לא זוהה בביטחון כאנגלית"],
            ["Bot walls & junk pages", "junk_page", "CAPTCHA, חומת תשלום, דפי שגיאה"],
            ["Spam & malformed text", "hashtag_soup, repetitive, low_alpha_ratio",
             "האשטגים, חזרתיות, שאריות בינאריות"],
            ["Duplicates & reposts", "duplicate, repost", "טקסט זהה שכבר התקבל או הדהוד"],
            ["Replies & flagged posts", "reply, sensitive", "תגובות אישיות ופוסטים עם אזהרה"],
            ["No extractable themes", "no_keywords", "KeyBERT לא מצא ביטוי מפתח בטוח"],
            ["Unreachable", "unreachable, processing_error",
             "קישורים מתים, פסקי זמן ורשומות פגומות"],
        ],
        "קיבוץ קודי הדחייה לקטגוריות תצוגה", widths=[4.0, 5.4, 6.4], font_size=9.5)


def chapter4(R):
    R.heading("פרק 4: עיבוד שפה טבעית וניתוח גרפים", 1, page_break=True)
    R.heading("4.1 מודל השיכון המשותף", 2)
    R.para(
        "המערכת משתמשת במקודד משפטים יחיד — all-MiniLM-L6-v2 (Wang et al., 2020), "
        "בגישת Sentence-BERT (Reimers & Gurevych, 2019) — הן לחילוץ ביטויי המפתח "
        "בשלב הקליטה והן לבניית הגרף הסמנטי בשלב הניתוח. כך הביטויים והקשתות "
        "\"חיים\" באותו מרחב שיכון בן 384 ממדים. MiniLM נבחר בשל גודלו הקטן, המאפשר "
        "הסקה מהירה על CPU.")

    R.heading("4.2 חילוץ ביטויי מפתח עם KeyBERT", 2)
    R.para(
        "KeyBERT מחשב את שיכון המסמך ואת שיכוני המועמדים (n-grams באורך 1–2 לאחר "
        "הסרת מילות עצירה), ומדרג את המועמדים לפי דמיון קוסינוס למסמך. כדי למנוע "
        "ביטויים כמעט זהים מופעל Maximal Marginal Relevance (Carbonell & Goldstein, "
        "1998) עם diversity = 0.6, המאזן בין רלוונטיות למסמך לבין שוני מן הביטויים "
        "שכבר נבחרו. נשמרים עד חמישה ביטויים בעלי ציון גבוה מ-0.35.")
    R.para(
        "בסריקות הרקע מופקות \"מילות באזז\" לכל פלטפורמה (engine/signals.py) מתוך "
        "ביטויי KeyBERT של הפוסטים עצמם, ללא הפעלה נוספת של המודל. מסנן נוסף "
        "(clean_buzzword) משאיר רק ביטויים של מילה או שתיים באותיות לטיניות קטנות, "
        "ומסיר מילות עצירה של sklearn ומילות מילוי שיחתיות (just, really, actually). "
        "כל ביטוי נספר פעם אחת לכל פוסט, כך שפוסט מילולי אחד אינו יכול לשלוט ברשימה, "
        "ונשמרות שמונה המילים השכיחות ביותר.")

    R.heading("4.3 ניתוח סנטימנט עם VADER", 2)
    R.para(
        "לצורך מדידת \"זיקת הקהילה\" לנושא נבחר VADER (Hutto & Gilbert, 2014) — "
        "מודל מבוסס-כללים ולקסיקון שדורג על-ידי בני אדם, הכולל כללים לשלילה, "
        "מעצימים, פיסוק ואותיות רישיות. VADER תוכנן לטקסט קצר ברשתות חברתיות, אינו "
        "דורש הורדת מודל ומנקד פוסט במיקרו-שניות, ולכן אינו מוסיף זמן מורגש לסריקה — "
        "בניגוד למודל Transformer, שהיה מכפיל את עלות ה-CPU לפוסט.")
    R.para(
        "לקסיקון כללי מפרש באופן שגוי אוצר מילים טכני: \"kill a zombie process\" "
        "אינו משפט זועם. לכן אופסה הוולנטיות של כ-50 מונחים טכניים (kill, zombie, "
        "deadlock, fatal, exploit, crash, master, slave ועוד). הניקוד מחושב על הכותרת "
        "ו-300 התווים הראשונים של הגוף, ומתקבל ציון compound בטווח [−1, 1]. הנושא "
        "מסווג לפי הכללים הבאים:")
    R.table(
        ["תווית", "תנאי (לפי סדר הבדיקה)"],
        [
            ["Mixed", "שיעור הפוסטים החיוביים וגם השליליים ≥ 0.30 — קודם לכל השאר"],
            ["Highly positive", "ממוצע compound ≥ 0.35"],
            ["Positive", "ממוצע compound ≥ 0.05"],
            ["Negative", "ממוצע compound ≤ −0.05"],
            ["Neutral", "בכל מקרה אחר"],
        ],
        "כללי סיווג זיקת הקהילה (engine/sentiment.py)", widths=[4.0, 11.0])
    R.para(
        "תווית Mixed קודמת לכל השאר, משום שדיון מקוטב אינו \"ניטרלי\" רק משום "
        "שהפוסטים החיוביים והשליליים מקזזים זה את זה בממוצע. פוסט בודד נחשב חיובי או "
        "שלילי כאשר ציונו חוצה ±0.05.")

    R.heading("4.4 בניית גרף הדמיון הסמנטי", 2)
    R.para(
        "המחלקה GraphBuilder (ui/graph_analyzer.py) הופכת את הקורפוס לגרף לא-מכוון "
        "ממושקל G = (V, E), שבו כל צומת הוא פוסט וכל קשת מחברת שני פוסטים שהשיכונים "
        "שלהם דומים מעל סף אדפטיבי. המודול אינו תלוי ב-Streamlit; המטמון הוא באחריות "
        "הקורא. הצינור כולל חמישה שלבים:")
    R.row_of_boxes([
        ("דגימה", "Top-25 לפלטפורמה"),
        ("שיכון", "MiniLM, נורמה 1"),
        ("קישור", "ספים מרובדים"),
        ("אשכול", "מודולריות CNM"),
        ("דירוג", "גשרים מתהווים"),
    ], fill=BAND_FILLS[0])
    R.caption(FIG_LABEL, "צינור ניתוח הגרף הסמנטי")
    R.bullets([
        "**דגימה מאוזנת:** 25 הפוסטים המובילים בכל פלטפורמה לפי ציון מגמה — שאילתה "
        "חסומה אחת לכל פלטפורמה מבטיחה איזון.",
        "**שיכון:** הכותרת ו-600 התווים הראשונים של הגוף (המקום שבו פוסט מציג את "
        "נושאו) מקודדים לווקטורים מנורמלים, כך שמכפלה פנימית שווה לדמיון קוסינוס.",
        "**קישור בספים מרובדים:** פוסטים מאותה פלטפורמה דומים זה לזה באופן שיטתי "
        "(סגנון, תמהיל נושאים, אורך), ולכן סף גלובלי יחיד היה ממלא את הגרף בקשתות "
        "פנים-פלטפורמיות לפני שקשת חוצת-פלטפורמות אחת הייתה עוברת. לכן מחושבים שני "
        "ספים נפרדים — לזוגות מאותה פלטפורמה ולזוגות חוצים.",
    ])
    R.formula("cutoff(S) = max( P97({ sim(i, j) : (i, j) ∈ S }), 0.35 )")
    R.para(
        "האחוזון ה-97 נבחר אמפירית כנקודה שבה מבנה האשכולות החזק ביותר (מודולריות "
        "של כ-0.5) מתקבל בלי שהגרף יירווה, ורצפת המוחלט 0.35 מונעת מקורפוס קטן או "
        "הומוגני \"לייצר\" קשתות בין פוסטים שאינם קשורים. מילות המפתח השמורות אינן "
        "יוצרות קשתות: גיוון ה-MMR הופך אותן לספציפיות למסמך, וגרף מבוסס מילות מפתח "
        "משותפות היה כמעט חסר קשתות. הן משמשות במקום זאת להסבר מדוע שני פוסטים "
        "מקושרים.")

    R.heading("4.5 זיהוי קהילות", 2)
    R.para(
        "הקהילות בגרף מזוהות באלגוריתם החמדני למקסום מודולריות של Clauset, Newman "
        "ו-Moore (2004), במימוש של NetworkX (Hagberg et al., 2008), תוך שימוש במשקלי "
        "הקשתות. המודולריות (Newman, 2006) מוגדרת:")
    R.formula("Q = (1 / 2m) · Σ_ij [ A_ij − k_i·k_j / 2m ] · δ(c_i, c_j)")
    R.para(
        "תוויות הקהילה מזינות את בונוס חציית הקהילות בדירוג הגשרים. בגרף מנוון "
        "(למשל ללא קשתות) המערכת מתדרדרת בחן ומדרגת ללא הבונוס.")

    R.heading("4.6 גשרים מתהווים ופערים מבניים", 2)
    R.para(
        "הבסיס התאורטי לזיהוי הגשרים הוא \"חוזקם של קשרים חלשים\" (Granovetter, "
        "1973) ותאוריית \"החורים המבניים\" (Burt, 1992): מידע חדשני זורם דווקא דרך "
        "קשרים המחברים קבוצות שאחרת אינן מחוברות. בהקשר של הפרויקט, גשר מתהווה הוא "
        "זוג פוסטים מקהילות רחוקות מבנית העוסקים באותו רעיון תוך שיתוף שכנים מועט — "
        "סימן לכך שאותו רעיון קיים בשני חלקים מנותקים של המערכת האקולוגית. "
        "detect_emerging_bridges (ui/graph_bridges.py) דורש שלושה תנאים, שכל אחד "
        "מהם שולל סוג אחר של חיובי-שגוי:")
    R.bullets([
        "**דמיון גבוה:** משקל הקשת חוצת-הפלטפורמות עובר את האחוזון ה-75 של משקלי "
        "הקשתות החוצות — סף המותאם לקורפוס וזונח את הקשתות החלשות שעשויות להיות "
        "מקריות.",
        "**שיבוץ נמוך (Weak tie):** מקדם Jaccard של שכונות שני הקצוות אינו עולה על "
        "0.25. ללא תנאי זה, \"גשרים\" היו בסך הכול הליבה הצפופה של אשכול אחד הפרוש "
        "על שתי פלטפורמות.",
        "**פער משלב:** הפלטפורמות רחוקות זו מזו על ציר קוד-לשיח של לפחות 0.35. "
        "התנאי מתיר כל זוג פרט ל-Dev.to ו-Hacker News (פער 0.25), הנושאים שניהם "
        "פרשנות בצורת מאמר לאותו קהל.",
    ], numbered=True)
    R.table(
        ["פלטפורמה", "מיקום על ציר קוד←שיח", "תווית משלב"],
        [
            ["GitHub", "0.00", "implementation — מימוש"],
            ["Dev.to", "0.40", "practitioner write-up — כתיבת מתרגלים"],
            ["Hacker News", "0.65", "industry debate — דיון תעשייתי"],
            ["Mastodon", "1.00", "public discussion — שיח ציבורי"],
        ],
        "מיקום הפלטפורמות על ציר המשלב (PLATFORM_REGISTER)", widths=[3.5, 4.5, 7.0])
    R.para(
        "הציר אינו מדרג פופולריות או איכות, אלא את אופי השיח — מן התוצר (קוד) ועד "
        "השיח הציבורי. הציון של כל גשר מועמד מחושב כמכפלה:")
    R.formula("score = sim(u, v) × (1 − J(N(u), N(v))) × gap(p_u, p_v) × bonus")
    R.para(
        "כאשר bonus = 1.25 אם שני הקצוות שייכים לקהילות שונות, ו-1.0 אחרת. נבחרה "
        "מכפלה ולא סכום ממושקל משום שכל גורם הוא תנאי הכרחי: בסכום, דמיון גבוה היה "
        "יכול לפצות על פער משלב אפסי. בונוס הקהילה הוא משקל רך ולא מסנן קשיח, משום "
        "שמקסום מודולריות חמדני הוא היוריסטיקה על גרפים קטנים. לבסוף מתבצעת בחירה "
        "בשני מעברים: לפי סדר הציון, עם תקרה של שני גשרים לכל זוג פלטפורמות "
        "(עיקרון דומה ל-MMR, שאחרת GitHub↔Mastodon עם פער 1.0 היה שולט בכל "
        "המקומות), ולאחר מכן השלמה מן המועמדים שנדחו, עד שישה גשרים בסך הכול.")
    R.para(
        "לכל גשר מופק הסבר בשפה טבעית ללא קריאה למודל שפה: \"המונחים המשותפים\" "
        "נגזרים מביטויי KeyBERT, מהכותרת ומפתיחת הגוף של שני הצמתים, לאחר הסרת "
        "מונחים המופיעים ביותר מרבע מצומתי הגרף (MAX_TERM_DOC_FREQ = 0.25) — מונחים "
        "אלה מתארים את הקורפוס כולו (בדרך כלל את נושא החיפוש עצמו) ואינם מסבירים "
        "דבר על הזוג. כך התצוגה אינה יכולה להיתקל במכסת API, להיכשל ברשת או לטעון "
        "דבר שהנתונים אינם תומכים בו.")
    R.table(
        ["רכיב", "פרמטר", "ערך"],
        [
            ["סף קשת", "percentile / min_similarity", "97 / 0.35"],
            ["סף גשר", "BRIDGE_SIMILARITY_PERCENTILE", "75"],
            ["שיבוץ חלש", "MAX_NEIGHBOUR_OVERLAP (Jaccard)", "0.25"],
            ["פער משלב", "MIN_REGISTER_GAP", "0.35"],
            ["בונוס קהילה", "cross_community", "×1.25"],
            ["גיוון", "MAX_PER_PLATFORM_PAIR / MAX_EMERGING_BRIDGES", "2 / 6"],
        ],
        "פרמטרי אלגוריתם זיהוי הגשרים", widths=[3.5, 7.5, 3.0])


def chapter5(R):
    R.heading("פרק 5: ממשק המשתמש ולוח המחוונים", 1, page_break=True)
    R.heading("5.1 מבנה שכבת התצוגה", 2)
    R.para(
        "שכבת התצוגה נבנתה ב-Streamlit ומחולקת למודולים לפי תחומי אחריות: app.py "
        "(אתחול, ניתוב, באנר הנושא ושורת מדדי KPI), styles.py (צבעים וגיליון סגנון), "
        "utils.py (רכיבי HTML ותרשימים משותפים), data.py (טוענים ממוטמנים ושירותי "
        "רקע), scan.py (ScanJob ברקע והתקדמותו), ו-views/ (דף הבית, שורת החיפוש "
        "ומודול לכל לשונית). הטוענים מוגדרים עם st.cache_data לחמש דקות ומבוטלים "
        "במפורש לאחר כל סריקה, כך שאינטראקציות עם רכיבים — כל אחת מהן הרצה מחדש של "
        "הסקריפט — אינן מתשאלות את בסיס הנתונים מחדש.")
    R.para(
        "צבעי הפלטפורמות (GitHub כחול, Hacker News כתום, Dev.to ירוק, Mastodon סגול) "
        "אחידים בכל התצוגות ונבדקו מול לקויות ראיית צבעים (מרחק ΔE מינימלי של 9.2 "
        "בהדמיית CVD). טקסט לעולם אינו נצבע בצבע סדרה.")

    R.heading("5.2 דף הבית ותדריך המודיעין", 2)
    R.para(
        "דף הבית מציג \"תדריך מודיעין\" (Intelligence Brief) שמוזן מעובד הרקע: "
        "טבלת מובילים של הנושאים שנסרקו, תובנות, מילות באזז, הפוסטים המובילים "
        "וכפתורי \"Deep Dive\" המפעילים סריקה מלאה של נושא. נושא מסומן כבעל תנופה "
        "חזקה (🔥) כאשר יש לו פוסטים תקינים בשלוש קהילות לפחות ו-12 פוסטים לפחות "
        "בסך הכול. ההחלטה מבוססת על ספירות ולא על ציוני מעורבות, משום שציוני ה-Z "
        "יחסיים לדגימה של כל נושא ואינם מאפשרים השוואה בין נושאים. בהפעלה הראשונה "
        "מוצג מסך Cold Start, ובזמן ההמתנה מוצגות כתבות מובילות מ-Dev.to.")
    R.screenshot_placeholder("דף הבית — תדריך המודיעין")
    R.caption(FIG_LABEL, "דף הבית: תדריך המודיעין של עובד הרקע")

    R.heading("5.3 לשונית 1: גלגל המערכת האקולוגית (Ecosystem Wheel)", 2)
    R.para(
        "הלשונית מציגה תרשים Sunburst של Plotly: הטבעת הפנימית מייצגת את הפלטפורמות "
        "והטבעת החיצונית את הפוסטים, וזווית כל פלח פרופורציונית לציון המגמה. התרשים "
        "נבנה ב-graph_objects, כך שלכל צומת מזהה יציב (root, platform:<name>, "
        "post:<id>) — פוסטים בעלי כותרות זהות נשארים נפרדים, ולחיצה מתורגמת חזרה "
        "לשורה בבסיס הנתונים. מאחר ש-st.plotly_chart אינו חושף את האירוע "
        "plotly_sunburstclick, פותח רכיב Streamlit מותאם (ui/components/wheel_click) "
        "המאזין ללחיצות ומעביר אותן לפאנל \"Trend Inspector\" הצדדי. מתחת לגלגל "
        "מוצגת טבלה מדורגת של כל הפוסטים בחיפוש.")
    R.screenshot_placeholder("לשונית Ecosystem Wheel")
    R.caption(FIG_LABEL, "לשונית גלגל המערכת האקולוגית ופאנל Trend Inspector")

    R.heading("5.4 לשונית 2: הייפ מול מציאות (Hype vs. Reality)", 2)
    R.para(
        "הלשונית משווה בין נפח הקוד (GitHub) לבין נפח השיח (Hacker News, Dev.to "
        "ו-Mastodon). מאחר שבבסיס הנתונים נשמרים לכל היותר 25 פוסטים לפלטפורמה, "
        "ספירת הפוסטים השמורים הייתה מציגה תמיד יחס קבוע של 1:3. לכן ההשוואה "
        "מבוססת על estimated_matches מטבלת scan_volume — הספירה של הפלטפורמה עצמה, "
        "מתוקנת לחפיפה ולשיעור הקבלה (סעיף 2.3.3). נפח נמדד משמש רק כאשר כל ארבע "
        "הפלטפורמות דיווחו עליו; אחרת מוצגות ספירות שמורות עם סימון מתאים. ספירות "
        "שהגיעו לתקרת הדיווח של ה-API מסומנות כ-\"לפחות\".")
    R.table(
        ["פסק דין", "תנאי", "פרשנות"],
        [
            ["Early Signal 🌱", "פחות מ-30 פריטים תואמים בסך הכול",
             "מוקדם מכדי לקבוע אם מדובר בהייפ או במציאות"],
            ["Backed by Code 🛠️", "נתח הקוד ≥ 60%", "הבונים מקדימים את הבאזז"],
            ["All Hype 📣", "נתח הקוד ≤ 30%", "שיח רב, מעט קוד מתוחזק"],
            ["Balanced ⚖️", "בכל מקרה אחר", "קוד ושיח צומחים יחד"],
        ],
        "רצועות פסק הדין בלשונית Hype vs. Reality", widths=[3.8, 5.2, 6.5])
    R.screenshot_placeholder("לשונית Hype vs. Reality")
    R.caption(FIG_LABEL, "לשונית הייפ מול מציאות")

    R.heading("5.5 לשונית 3: חיבורים (Connections)", 2)
    R.para(
        "הלשונית מציגה כרטיס ראיות לכל גשר מתהווה שזוהה (פרק 4.6): שני הפוסטים "
        "וקישוריהם, המשלב של כל קהילה, ההסבר בשפה טבעית ושלושה אריחי ראיות — עוצמת "
        "ההתאמה (Very strong מ-0.70, Strong מ-0.55, אחרת Moderate), רוחב הגשר לפי "
        "פער המשלב (Very wide מ-0.9, Wide מ-0.6) ומידת ההיעדר של שכנים משותפים. "
        "ביטויי המפתח של כל צד מוצגים כ\"שבבים\", והמונחים המשותפים מודגשים. כל "
        "ההסברים נגזרים ממדידות הגרף עצמו, ללא קריאה למודל.")
    R.screenshot_placeholder("לשונית Connections")
    R.caption(FIG_LABEL, "לשונית החיבורים: כרטיסי ראיות לגשרים מתהווים")

    R.heading("5.6 לשונית 4: מסנן שומר הסף (Gatekeeper's Filter)", 2)
    R.para(
        "הלשונית הופכת את Gatekeeper לשקוף: היא מציגה שורת KPI (מועמדים שנבדקו, "
        "שהתקבלו ושנחסמו), תרשים משפך המראה כיצד מצטמצם מאגר המועמדים עד לפוסטים "
        "השמורים, ותרשים עמודות אופקי של הדחיות לפי הקטגוריות שבטבלה 4. מאחר שכל "
        "מועמד שנבדק מסתיים במצב אחד בדיוק — התקבל, נדחה עם סיבה או לא היה נגיש — "
        "המשפך מתאזן תמיד מול מה שנשמר. הלשונית מציגה גם את חמשת השלבים לפי סדר "
        "הביצוע ואת פירוט הדחיות לכל פלטפורמה.")
    R.screenshot_placeholder("לשונית Gatekeeper's Filter")
    R.caption(FIG_LABEL, "לשונית מסנן שומר הסף: משפך וסיבות דחייה")

    R.heading("5.7 חוויית הסריקה", 2)
    R.para(
        "סריקה אינטראקטיבית רצה כ-ScanJob בתהליכון רקע, ופאנל סטטוס חי מציג את "
        "שלבי run_cycle (loading, planning, collecting, platform_done, scoring) "
        "ואת מספר הפוסטים שהתקבלו מכל פלטפורמה. בזמן סריקה הממשק ננעל, כדי שסשנים "
        "שונים לא יחליפו זה את הקורפוס של זה. סריקה שלא מצאה די תוצאות משאירה את "
        "הקורפוס הקודם על כנו ומציגה הודעה מוסברת (נפח נמוך, או פחות משלוש "
        "פלטפורמות).")


def chapter6(R):
    R.heading("פרק 6: סיכום, מסקנות ועבודה עתידית", 1, page_break=True)
    R.heading("6.1 סיכום", 2)
    R.para(
        "בפרויקט תוכננה ומומשה מערכת TrendAnalyzer Pro, האוספת מידע מארבע קהילות "
        "מפתחים הטרוגניות, מסננת ממנו רעש באופן שקוף, מנרמלת את מדדי המעורבות "
        "ומפיקה תובנות ניתנות להסבר: התפלגות תשומת הלב, היחס בין קוד לשיח, וגשרים "
        "סמנטיים בין קהילות. המערכת בנויה בארכיטקטורה תלת-שכבתית ברורה, כאשר "
        "המנוע אינו תלוי בממשק, שכבת הנתונים מבודדת מאחורי DAO, וכל הקבועים מרוכזים "
        "בקובץ תצורה יחיד.")

    R.heading("6.2 מסקנות", 2)
    R.bullets([
        "**סדר השלבים קובע את הביצועים:** הצבת הבדיקות המבניות הזולות לפני "
        "langdetect ו-KeyBERT מבטיחה שהמודלים הכבדים מעבדים רק חלק קטן מן הקלט.",
        "**שקיפות בונה אמון:** חשבונאות משפך מאוזנת, שבה כל מועמד נספר בדיוק פעם "
        "אחת, מאפשרת למשתמש להבין מדוע התקבלו התוצאות שלפניו — ולמפתח לכייל את "
        "הספים.",
        "**נרמול לכל פלטפורמה הכרחי:** ללא Z-score לוגריתמי לכל פלטפורמה וגיזום "
        "מרובד, GitHub (שבו ספירות הכוכבים גבוהות בסדרי גודל) היה שולט בכל התצוגות.",
        "**ספים אדפטיביים ומרובדים עדיפים על ספים קבועים:** אחוזונים המחושבים על "
        "הקורפוס, בנפרד לזוגות פנים-פלטפורמיים וחוצי-פלטפורמות, מאפשרים לקשתות "
        "חוצות להופיע בכלל.",
        "**מכפלה של תנאים הכרחיים:** ניסוח ציון הגשר כמכפלה מונע מגורם חזק אחד לפצות "
        "על היעדרו של אחר, ותקרת הגיוון מונעת מזוג פלטפורמות אחד לשלוט בתוצאות.",
    ])

    R.heading("6.3 אתגרים בארכיטקטורה האסינכרונית", 2)
    R.para("במהלך הפיתוח התגלו כמה אתגרים מהותיים, ולכל אחד נמצא פתרון מתועד:")
    R.table(
        ["אתגר", "תיאור", "פתרון"],
        [
            ["שילוב asyncio עם Streamlit",
             "Streamlit מריץ סקריפטים בתהליכון ללא לולאת אירועים, ולעיתים מותקנת לולאה "
             "על-ידי המארח", "_run_coroutine: asyncio.run או לולאה פרטית בתהליכון עובד"],
            ["פלטפורמה תקועה", "בקשה אחת שאינה מסתיימת עלולה להחזיק את הסריקה כולה",
             "wait_for עם 150 שניות לכל פלטפורמה ו-15 שניות לכל בקשה"],
            ["תקלה בפריט בודד", "חריגה בפוסט אחד ב-gather מבטלת את כל הגל",
             "return_exceptions=True ורישום processing_error"],
            ["מגבלות קצב", "GitHub (60/שעה ללא אסימון) ו-Dev.to מגבילים בקשות",
             "Semaphore לכל מארח, שמירת מכסה, CDN לפני API"],
            ["בטיחות מודלים בתהליכונים", "KeyBERT, הטוקנייזר ו-MPS אינם בטוחים "
             "לקריאות מקביליות", "Singleton, מנעולים, והצמדה ל-CPU"],
            ["כותב וקוראים בו-זמנית", "סריקה כותבת בזמן שהממשק ועובד הרקע קוראים",
             "SQLite WAL, busy timeout, חיבורים קצרים, בסיס נתונים זמני לרקע"],
            ["ביטול באמצע סריקה", "המשתמש עוזב או לוחץ על כפתור בזמן איסוף",
             "Event נבדק כל 0.5 שניות; ביטול לפני כל כתיבה"],
            ["Cold Start", "דף הבית ריק עד שהעובד משלים את המחזור הראשון",
             "מסך טעינה חוסם ותזמון השלמה לנושאים חסרים"],
        ],
        "אתגרים מרכזיים בארכיטקטורה האסינכרונית ופתרונותיהם",
        widths=[3.4, 6.0, 6.4], font_size=9)

    R.heading("6.4 מגבלות", 2)
    R.bullets([
        "המערכת מוגבלת לתוכן באנגלית, הן בשל אוצר המילים של מסנן הרלוונטיות והן "
        "בשל מודלי השפה.",
        "Mastodon נאסף משרת יחיד (mastodon.social), ולכן אינו מייצג את כל ה-Fediverse.",
        "VADER הוא מודל לקסיקלי; גם לאחר נטרול מונחים טכניים הוא אינו מבין אירוניה "
        "או הקשר מורכב.",
        "הערכת הנפח מבוססת על ספירות שהפלטפורמות מדווחות, שחלקן מוגבלות בתקרה "
        "(למשל עמוד של 1,000 ב-Dev.to).",
        "חלון הזמן קצר (7–14 ימים), כך שהמערכת מתארת \"תמונת מצב\" ולא סדרה עתית "
        "ארוכת טווח.",
    ])

    R.heading("6.5 כיווני המשך", 2)
    R.bullets([
        "**ניתוח עתי:** שמירת היסטוריה ארוכת טווח של live_trends וזיהוי פרצי פעילות "
        "בגישת Kleinberg, כדי להבחין בין מגמה עולה למגמה דועכת.",
        "**פלטפורמות נוספות:** Reddit, Stack Overflow, arXiv ו-Bluesky, בזכות "
        "תבנית BaseCollector המאפשרת להוסיף אוסף על-ידי מימוש מתודה אחת.",
        "**תמיכה רב-לשונית:** מקודד רב-לשוני (למשל multilingual MiniLM) ואוצר מילים "
        "לרלוונטיות בכמה שפות.",
        "**סנטימנט מבוסס-הקשר:** מודל Transformer קל המכוונן לטקסט טכני, שירוץ רק "
        "על הפוסטים שהתקבלו כדי לשמור על זמן הסריקה.",
        "**מערכת תורים מבוזרת:** העברת האיסוף לעובדים נפרדים (למשל Celery או Redis "
        "Streams) ומעבר ל-PostgreSQL לשם הרחבה מעבר למכונה אחת.",
        "**הערכה כמותית:** בניית מערך נתונים מתויג לאמידת הדיוק והכיסוי של "
        "Gatekeeper ושל זיהוי הגשרים מול שיפוט אנושי.",
    ])


def references(R):
    R.heading("רשימת מקורות", 1, page_break=True)
    refs = [
        "Burt, R. S. (1992). Structural holes: The social structure of competition. "
        "Harvard University Press.",
        "Carbonell, J., & Goldstein, J. (1998). The use of MMR, diversity-based reranking "
        "for reordering documents and producing summaries. In Proceedings of the 21st "
        "Annual International ACM SIGIR Conference (pp. 335-336). ACM.",
        "Clauset, A., Newman, M. E. J., & Moore, C. (2004). Finding community structure in "
        "very large networks. Physical Review E, 70(6), 066111.",
        "Fowler, M. (2002). Patterns of enterprise application architecture. "
        "Addison-Wesley.",
        "Granovetter, M. S. (1973). The strength of weak ties. American Journal of "
        "Sociology, 78(6), 1360-1380.",
        "Grootendorst, M. (2020). KeyBERT: Minimal keyword extraction with BERT "
        "[Computer software]. Zenodo. https://doi.org/10.5281/zenodo.4461265",
        "Hagberg, A. A., Schult, D. A., & Swart, P. J. (2008). Exploring network "
        "structure, dynamics, and function using NetworkX. In Proceedings of the 7th "
        "Python in Science Conference (SciPy 2008) (pp. 11-15).",
        "Heaps, H. S. (1978). Information retrieval: Computational and theoretical "
        "aspects. Academic Press.",
        "Hutto, C. J., & Gilbert, E. (2014). VADER: A parsimonious rule-based model for "
        "sentiment analysis of social media text. In Proceedings of the Eighth "
        "International AAAI Conference on Weblogs and Social Media (ICWSM-14).",
        "Kleinberg, J. (2003). Bursty and hierarchical structure in streams. Data Mining "
        "and Knowledge Discovery, 7(4), 373-397.",
        "Mathioudakis, M., & Koudas, N. (2010). TwitterMonitor: Trend detection over the "
        "Twitter stream. In Proceedings of the 2010 ACM SIGMOD International Conference "
        "on Management of Data (pp. 1155-1158). ACM.",
        "Newman, M. E. J. (2006). Modularity and community structure in networks. "
        "Proceedings of the National Academy of Sciences, 103(23), 8577-8582.",
        "Pedregosa, F., Varoquaux, G., Gramfort, A., et al. (2011). Scikit-learn: Machine "
        "learning in Python. Journal of Machine Learning Research, 12, 2825-2830.",
        "Reimers, N., & Gurevych, I. (2019). Sentence-BERT: Sentence embeddings using "
        "Siamese BERT-networks. In Proceedings of EMNLP-IJCNLP 2019 (pp. 3982-3992). ACL.",
        "Shuyo, N. (2010). Language detection library for Java [Computer software]. "
        "https://github.com/shuyo/language-detection",
        "Wang, W., Wei, F., Dong, L., Bao, H., Yang, N., & Zhou, M. (2020). MiniLM: Deep "
        "self-attention distillation for task-agnostic compression of pre-trained "
        "transformers. Advances in Neural Information Processing Systems, 33, 5776-5788.",
        "GitHub, Inc. (2026). GitHub REST API documentation: Search. "
        "https://docs.github.com/en/rest/search",
        "Hacker News & Algolia. (2026). HN Search API. https://hn.algolia.com/api",
        "Forem. (2026). Forem (Dev.to) API documentation. "
        "https://developers.forem.com/api",
        "Mastodon gGmbH. (2026). Mastodon API documentation: Timelines. "
        "https://docs.joinmastodon.org/methods/timelines/",
        "Python Software Foundation. (2026). asyncio — Asynchronous I/O. "
        "https://docs.python.org/3/library/asyncio.html",
        "SQLite Consortium. (2026). Write-ahead logging. https://www.sqlite.org/wal.html",
        "Snowflake Inc. (2026). Streamlit documentation. https://docs.streamlit.io",
    ]
    for i, ref in enumerate(refs, 1):
        p = R.english(f"[{i}]  {ref}", size=11)
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)


def build(R):
    abstract(R)
    chapter1(R)
    chapter2(R)
    chapter3(R)
    chapter4(R)
    chapter5(R)
    chapter6(R)
    references(R)


def main():
    # Pass 1 collects headings and captions, so the pre-rendered TOC and lists of
    # figures/tables (shown before Word refreshes the fields) are complete.
    probe = Report()
    build(probe)

    R = Report()
    title_page(R)
    R.headings, R.figures, R.tables = probe.headings, probe.figures, probe.tables
    front_matter(R)
    R.headings, R.figures, R.tables = [], [], []
    R.fig_no = R.tab_no = 0
    build(R)
    R.save(OUTPUT_PATH)
    print(f"Saved {OUTPUT_PATH}")
    print(f"  {len(R.headings)} headings, {len(R.figures)} figures, {len(R.tables)} tables")


if __name__ == "__main__":
    main()
